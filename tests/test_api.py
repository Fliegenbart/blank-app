"""
API Tests for Brand Engine.

These tests can be run with pytest:
    pytest tests/test_api.py -v

Note: Requires running services (docker compose up)
"""

import pytest
import httpx
import os

# Base URL for API
BASE_URL = os.getenv("API_URL", "http://localhost:8000")


@pytest.fixture
def auth_token():
    """Get an auth token for testing."""
    import uuid
    unique_id = str(uuid.uuid4())[:8]
    user = {
        "email": f"brand_test_{unique_id}@example.com",
        "password": "testpassword123",
    }
    httpx.post(f"{BASE_URL}/api/v1/auth/register", json=user)
    response = httpx.post(
        f"{BASE_URL}/api/v1/auth/login",
        data={"username": user["email"], "password": user["password"]},
    )
    return response.json()["access_token"]


class TestHealthCheck:
    """Test health check endpoint."""

    def test_health_check(self):
        """Test that health check returns healthy status."""
        response = httpx.get(f"{BASE_URL}/health")
        assert response.status_code == 200
        data = response.json()
        assert data["status"] == "healthy"


class TestAuth:
    """Test authentication endpoints."""

    @pytest.fixture
    def test_user(self):
        """Create a test user credentials."""
        import uuid
        unique_id = str(uuid.uuid4())[:8]
        return {
            "email": f"test_{unique_id}@example.com",
            "password": "testpassword123",
            "full_name": "Test User",
        }

    def test_register_user(self, test_user):
        """Test user registration."""
        response = httpx.post(
            f"{BASE_URL}/api/v1/auth/register",
            json=test_user,
        )
        assert response.status_code == 201
        data = response.json()
        assert data["email"] == test_user["email"]
        assert "id" in data

    def test_login_user(self, test_user):
        """Test user login."""
        # First register
        httpx.post(f"{BASE_URL}/api/v1/auth/register", json=test_user)

        # Then login
        response = httpx.post(
            f"{BASE_URL}/api/v1/auth/login",
            data={
                "username": test_user["email"],
                "password": test_user["password"],
            },
        )
        assert response.status_code == 200
        data = response.json()
        assert "access_token" in data
        assert data["token_type"] == "bearer"

    def test_get_current_user(self, test_user):
        """Test getting current user info."""
        # Register and login
        httpx.post(f"{BASE_URL}/api/v1/auth/register", json=test_user)
        login_response = httpx.post(
            f"{BASE_URL}/api/v1/auth/login",
            data={
                "username": test_user["email"],
                "password": test_user["password"],
            },
        )
        token = login_response.json()["access_token"]

        # Get current user
        response = httpx.get(
            f"{BASE_URL}/api/v1/users/me",
            headers={"Authorization": f"Bearer {token}"},
        )
        assert response.status_code == 200
        data = response.json()
        assert data["email"] == test_user["email"]


class TestBrands:
    """Test brand CRUD operations."""

    def test_create_brand(self, auth_token):
        """Test creating a brand."""
        import uuid
        unique_id = str(uuid.uuid4())[:8]

        response = httpx.post(
            f"{BASE_URL}/api/v1/brands",
            headers={"Authorization": f"Bearer {auth_token}"},
            json={
                "name": "Test Brand",
                "slug": f"test-brand-{unique_id}",
                "description": "A test brand",
            },
        )
        assert response.status_code == 201
        data = response.json()
        assert data["name"] == "Test Brand"
        assert "id" in data

    def test_list_brands(self, auth_token):
        """Test listing brands."""
        response = httpx.get(
            f"{BASE_URL}/api/v1/brands",
            headers={"Authorization": f"Bearer {auth_token}"},
        )
        assert response.status_code == 200
        data = response.json()
        assert isinstance(data, list)

    def test_get_brand(self, auth_token):
        """Test getting a specific brand."""
        import uuid
        unique_id = str(uuid.uuid4())[:8]

        # Create brand
        create_response = httpx.post(
            f"{BASE_URL}/api/v1/brands",
            headers={"Authorization": f"Bearer {auth_token}"},
            json={
                "name": "Get Test Brand",
                "slug": f"get-test-brand-{unique_id}",
            },
        )
        brand_id = create_response.json()["id"]

        # Get brand
        response = httpx.get(
            f"{BASE_URL}/api/v1/brands/{brand_id}",
            headers={"Authorization": f"Bearer {auth_token}"},
        )
        assert response.status_code == 200
        data = response.json()
        assert data["id"] == brand_id
        assert data["name"] == "Get Test Brand"

    def test_update_brand_member_role(self, auth_token):
        """Test updating a brand member role."""
        import uuid

        # Create second user
        member_email = f"member_{uuid.uuid4()}@example.com"
        member_password = "testpassword123"
        httpx.post(
            f"{BASE_URL}/api/v1/auth/register",
            json={"email": member_email, "password": member_password},
        )

        # Create brand
        create_response = httpx.post(
            f"{BASE_URL}/api/v1/brands",
            headers={"Authorization": f"Bearer {auth_token}"},
            json={
                "name": "Members Brand",
                "slug": f"members-brand-{str(uuid.uuid4())[:8]}",
            },
        )
        brand_id = create_response.json()["id"]

        # Add member
        add_response = httpx.post(
            f"{BASE_URL}/api/v1/brands/{brand_id}/members",
            headers={"Authorization": f"Bearer {auth_token}"},
            json={"user_email": member_email, "role": "viewer"},
        )
        assert add_response.status_code == 201
        member_id = add_response.json()["id"]

        # Update role
        update_response = httpx.put(
            f"{BASE_URL}/api/v1/brands/{brand_id}/members/{member_id}",
            headers={"Authorization": f"Bearer {auth_token}"},
            json={"role": "editor"},
        )
        assert update_response.status_code == 200
        assert update_response.json()["role"] == "editor"


class TestOutputs:
    """Test output generation and deletion."""

    def _create_sample_pptx(self, path: str):
        """Create a simple PPTX for upload."""
        from pptx import Presentation

        prs = Presentation()
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = "Sample Brand"
        slide.placeholders[1].text = "Sample content for testing"
        prs.save(path)

    def _poll_job(self, token: str, job_id: str, timeout: int = 60):
        import time

        headers = {"Authorization": f"Bearer {token}"}
        start = time.time()
        while time.time() - start < timeout:
            resp = httpx.get(f"{BASE_URL}/api/v1/jobs/{job_id}", headers=headers)
            resp.raise_for_status()
            status = resp.json()["status"]
            if status in ("completed", "failed"):
                return status
            time.sleep(2)
        return "timeout"

    def test_generate_and_delete_output(self):
        """Generate a website output and delete it."""
        import uuid
        import tempfile
        import os

        # Register and login
        email = f"output_{uuid.uuid4()}@example.com"
        password = "testpassword123"
        httpx.post(
            f"{BASE_URL}/api/v1/auth/register",
            json={"email": email, "password": password},
        )
        login_response = httpx.post(
            f"{BASE_URL}/api/v1/auth/login",
            data={"username": email, "password": password},
        )
        token = login_response.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}

        # Create brand
        create_response = httpx.post(
            f"{BASE_URL}/api/v1/brands",
            headers=headers,
            json={
                "name": "Output Brand",
                "slug": f"output-brand-{str(uuid.uuid4())[:8]}",
            },
        )
        brand_id = create_response.json()["id"]

        # Upload sample pptx
        with tempfile.TemporaryDirectory() as tmpdir:
            sample_path = os.path.join(tmpdir, "sample.pptx")
            self._create_sample_pptx(sample_path)
            with open(sample_path, "rb") as f:
                files = {
                    "file": (
                        "sample.pptx",
                        f,
                        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    )
                }
                upload_response = httpx.post(
                    f"{BASE_URL}/api/v1/brands/{brand_id}/uploads",
                    headers=headers,
                    files=files,
                )
        upload_id = upload_response.json()["id"]

        # Analyze upload
        analyze_response = httpx.post(
            f"{BASE_URL}/api/v1/brands/{brand_id}/analyze",
            headers=headers,
            params={"upload_id": upload_id},
        )
        analyze_job_id = analyze_response.json()["id"]
        assert self._poll_job(token, analyze_job_id) == "completed"

        # Generate website
        gen_response = httpx.post(
            f"{BASE_URL}/api/v1/brands/{brand_id}/generate/website",
            headers=headers,
            json={"topic": "Launch", "pages": [{"slug": "index", "title": "Home", "sections": ["hero"]}]},
        )
        gen_job_id = gen_response.json()["id"]
        assert self._poll_job(token, gen_job_id, timeout=90) == "completed"

        # List outputs
        outputs_response = httpx.get(
            f"{BASE_URL}/api/v1/brands/{brand_id}/outputs",
            headers=headers,
        )
        outputs = outputs_response.json()
        assert outputs
        output_id = outputs[0]["id"]

        # Delete output
        delete_response = httpx.delete(
            f"{BASE_URL}/api/v1/outputs/{output_id}",
            headers=headers,
        )
        assert delete_response.status_code == 204


class TestFiles:
    """Test file download endpoint permissions."""

    def test_files_endpoint_rejects_unknown(self, auth_token):
        response = httpx.get(
            f"{BASE_URL}/api/v1/files/unknown/path",
            headers={"Authorization": f"Bearer {auth_token}"},
        )
        assert response.status_code in (403, 404)


class TestAnalyzerUnits:
    """Unit tests for analyzers (don't require running services)."""

    def test_color_hex_to_hsl(self):
        """Test hex to HSL conversion."""
        from apps.worker.app.analyzers.color_analyzer import ColorAnalyzer

        analyzer = ColorAnalyzer()

        # Test pure red
        h, s, l = analyzer._hex_to_hsl("#FF0000")
        assert abs(h - 0) < 1  # Hue ~0
        assert abs(s - 1) < 0.01  # Full saturation
        assert abs(l - 0.5) < 0.01  # Mid lightness

        # Test white
        h, s, l = analyzer._hex_to_hsl("#FFFFFF")
        assert abs(l - 1) < 0.01  # Full lightness

        # Test black
        h, s, l = analyzer._hex_to_hsl("#000000")
        assert abs(l - 0) < 0.01  # No lightness

    def test_color_name_detection(self):
        """Test color name detection."""
        from apps.worker.app.analyzers.color_analyzer import ColorAnalyzer

        analyzer = ColorAnalyzer()

        assert analyzer._get_color_name("#FF0000") == "Red"
        assert analyzer._get_color_name("#0000FF") == "Blue"
        assert analyzer._get_color_name("#00FF00") == "Green"
        assert analyzer._get_color_name("#FFFFFF") == "White"
        assert analyzer._get_color_name("#000000") == "Black"


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
