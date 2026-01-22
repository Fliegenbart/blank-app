#!/usr/bin/env python3
"""
Generate synthetic sample files for testing the Brand Engine.
"""

import os
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Try to import reportlab for PDF generation
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False
    print("Warning: reportlab not installed, PDF generation will be skipped")


def generate_sample_pptx(output_path: str = "sample.pptx"):
    """Generate a sample PowerPoint presentation."""
    print(f"Generating {output_path}...")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title slide
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)

    # Add title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = "Acme Corporation"
    title_para.font.name = "Arial"
    title_para.font.size = Pt(60)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(0, 102, 204)  # Blue
    title_para.alignment = PP_ALIGN.CENTER

    # Add subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(11), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.text = "Innovation Through Excellence"
    subtitle_para.font.name = "Arial"
    subtitle_para.font.size = Pt(28)
    subtitle_para.font.color.rgb = RGBColor(102, 102, 102)  # Gray
    subtitle_para.alignment = PP_ALIGN.CENTER

    # Slide 2: Content slide
    slide2 = prs.slides.add_slide(slide_layout)

    # Add heading
    heading_box = slide2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
    heading_frame = heading_box.text_frame
    heading_para = heading_frame.paragraphs[0]
    heading_para.text = "Our Core Values"
    heading_para.font.name = "Arial"
    heading_para.font.size = Pt(44)
    heading_para.font.bold = True
    heading_para.font.color.rgb = RGBColor(0, 102, 204)

    # Add bullet points
    content_box = slide2.shapes.add_textbox(Inches(0.75), Inches(1.75), Inches(11), Inches(5))
    content_frame = content_box.text_frame
    content_frame.word_wrap = True

    values = [
        ("Innovation", "We embrace new ideas and technologies"),
        ("Quality", "We deliver excellence in everything we do"),
        ("Integrity", "We act with honesty and transparency"),
        ("Collaboration", "We work together to achieve more"),
    ]

    for i, (value, description) in enumerate(values):
        if i == 0:
            p = content_frame.paragraphs[0]
        else:
            p = content_frame.add_paragraph()

        p.text = f"• {value}: {description}"
        p.font.name = "Arial"
        p.font.size = Pt(24)
        p.font.color.rgb = RGBColor(51, 51, 51)
        p.space_after = Pt(20)

    # Slide 3: Contact slide
    slide3 = prs.slides.add_slide(slide_layout)

    # Add heading
    contact_heading = slide3.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1))
    ch_frame = contact_heading.text_frame
    ch_para = ch_frame.paragraphs[0]
    ch_para.text = "Get in Touch"
    ch_para.font.name = "Arial"
    ch_para.font.size = Pt(48)
    ch_para.font.bold = True
    ch_para.font.color.rgb = RGBColor(0, 102, 204)
    ch_para.alignment = PP_ALIGN.CENTER

    # Add contact info
    contact_box = slide3.shapes.add_textbox(Inches(1), Inches(4), Inches(11), Inches(2))
    c_frame = contact_box.text_frame
    c_para = c_frame.paragraphs[0]
    c_para.text = "www.acme-corp.example | hello@acme-corp.example"
    c_para.font.name = "Arial"
    c_para.font.size = Pt(24)
    c_para.font.color.rgb = RGBColor(255, 102, 0)  # Orange accent
    c_para.alignment = PP_ALIGN.CENTER

    # Save
    prs.save(output_path)
    print(f"Created {output_path}")


def generate_sample_pdf(output_path: str = "sample.pdf"):
    """Generate a sample PDF document."""
    if not HAS_REPORTLAB:
        print("Skipping PDF generation (reportlab not installed)")
        return

    print(f"Generating {output_path}...")

    c = canvas.Canvas(output_path, pagesize=letter)
    width, height = letter

    # Colors
    primary_blue = HexColor("#0066CC")
    dark_gray = HexColor("#333333")
    orange = HexColor("#FF6600")

    # Page 1: Cover
    c.setFont("Helvetica-Bold", 48)
    c.setFillColor(primary_blue)
    c.drawCentredString(width / 2, height - 250, "Acme Corporation")

    c.setFont("Helvetica", 24)
    c.setFillColor(dark_gray)
    c.drawCentredString(width / 2, height - 300, "Brand Guidelines 2025")

    c.setFont("Helvetica", 16)
    c.drawCentredString(width / 2, height - 400, "Innovation Through Excellence")

    c.showPage()

    # Page 2: Content
    c.setFont("Helvetica-Bold", 36)
    c.setFillColor(primary_blue)
    c.drawString(72, height - 72, "About Us")

    c.setFont("Helvetica", 12)
    c.setFillColor(dark_gray)

    text_content = """
    Acme Corporation is a leading provider of innovative solutions. We believe in
    delivering excellence through technology and collaboration.

    Our Mission:
    To empower businesses with cutting-edge solutions that drive growth and success.

    Our Values:
    • Innovation - We embrace new ideas and technologies
    • Quality - We deliver excellence in everything we do
    • Integrity - We act with honesty and transparency
    • Collaboration - We work together to achieve more

    Contact us at hello@acme-corp.example to learn more about how we can help your
    business succeed in today's competitive landscape.
    """.strip()

    y_position = height - 120
    for line in text_content.split('\n'):
        c.drawString(72, y_position, line.strip())
        y_position -= 18

    # Footer
    c.setFont("Helvetica", 10)
    c.setFillColor(orange)
    c.drawCentredString(width / 2, 50, "www.acme-corp.example")

    c.showPage()
    c.save()
    print(f"Created {output_path}")


def main():
    """Generate all sample files."""
    # Ensure we're in the samples directory or create files there
    script_dir = os.path.dirname(os.path.abspath(__file__))
    os.chdir(script_dir)

    print("Generating sample files for Brand Engine testing...\n")

    generate_sample_pptx("sample.pptx")
    generate_sample_pdf("sample.pdf")

    print("\nDone! Sample files are ready for testing.")
    print("\nTo test the pipeline:")
    print("1. Start services: cd ../infra && docker compose up -d")
    print("2. Run the API tests or use curl commands from README.md")


if __name__ == "__main__":
    main()
