"""Presentation generator for theme assets."""
import io
from typing import Dict, Any, List
import structlog

from app.providers.content_provider import get_content_provider

logger = structlog.get_logger()


class PresentationGenerator:
    """Generates PowerPoint presentation using python-pptx."""

    def __init__(self):
        self.content_provider = get_content_provider()

    def generate(self, context: Dict[str, Any]) -> bytes:
        """Generate a PPTX presentation."""
        profile = context.get("profile", {})
        theme = context.get("theme", {})

        logger.info("Generating presentation", theme=theme.get("name"))

        # Generate content using AI
        content = self._generate_content(profile, theme)

        # Create PPTX
        pptx_data = self._create_pptx(profile, theme, content)

        return pptx_data

    def _generate_content(self, profile: Dict[str, Any], theme: Dict[str, Any]) -> Dict[str, Any]:
        """Generate presentation content using AI."""
        theme_name = theme.get("name", "Event")
        theme_desc = theme.get("description", "")
        brand_name = profile.get("identity", {}).get("name", "Brand")

        doc_context = ""
        for doc in theme.get("documents", []):
            if doc.get("text"):
                doc_context += f"\n{doc['name']}:\n{doc['text'][:2000]}\n"

        prompt = f"""Create a professional presentation for "{theme_name}".

Brand: {brand_name}
Theme Description: {theme_desc}
{f"Additional Context: {doc_context}" if doc_context else ""}

Generate JSON with 8-12 slides:
{{
  "title": "Presentation title",
  "subtitle": "Subtitle or tagline",
  "slides": [
    {{
      "type": "title",
      "title": "Main title",
      "subtitle": "Subtitle"
    }},
    {{
      "type": "agenda",
      "title": "Agenda",
      "items": ["Topic 1", "Topic 2", "Topic 3"]
    }},
    {{
      "type": "content",
      "title": "Section Title",
      "content": "Main content or description",
      "bullets": ["Point 1", "Point 2", "Point 3"]
    }},
    {{
      "type": "two_column",
      "title": "Comparison",
      "left_title": "Column 1",
      "left_content": ["Item 1", "Item 2"],
      "right_title": "Column 2",
      "right_content": ["Item 1", "Item 2"]
    }},
    {{
      "type": "quote",
      "quote": "Inspiring quote",
      "author": "Author name"
    }},
    {{
      "type": "closing",
      "title": "Thank You",
      "subtitle": "Contact info or next steps"
    }}
  ]
}}

Include: title slide, agenda, 4-6 content slides, and closing."""

        try:
            result = self.content_provider.generate_json(prompt)
            return result
        except Exception as e:
            logger.error("Content generation failed, using defaults", error=str(e))
            return self._default_content(theme_name, brand_name)

    def _default_content(self, theme_name: str, brand_name: str) -> Dict[str, Any]:
        """Return default content."""
        return {
            "title": theme_name,
            "subtitle": f"Presented by {brand_name}",
            "slides": [
                {"type": "title", "title": theme_name, "subtitle": f"Presented by {brand_name}"},
                {"type": "agenda", "title": "Agenda", "items": ["Introduction", "Overview", "Key Points", "Next Steps"]},
                {"type": "content", "title": "Introduction", "content": f"Welcome to {theme_name}", "bullets": ["Background", "Objectives", "Scope"]},
                {"type": "content", "title": "Overview", "content": "Key highlights and information", "bullets": ["Point one", "Point two", "Point three"]},
                {"type": "content", "title": "Key Points", "content": "Important details", "bullets": ["Detail one", "Detail two", "Detail three"]},
                {"type": "closing", "title": "Thank You", "subtitle": "Questions?"}
            ]
        }

    def _create_pptx(self, profile: Dict[str, Any], theme: Dict[str, Any], content: Dict[str, Any]) -> bytes:
        """Create PPTX file using python-pptx."""
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RgbColor
            from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        except ImportError:
            logger.warning("python-pptx not installed, generating placeholder")
            return self._create_placeholder_pptx(content)

        # Create presentation
        prs = Presentation()
        prs.slide_width = Inches(13.333)  # 16:9
        prs.slide_height = Inches(7.5)

        # Get brand colors
        colors = profile.get("colors", {})
        primary_hex = colors.get("primary", {}).get("hex", "#0066CC").lstrip("#")
        try:
            primary_color = RgbColor(
                int(primary_hex[0:2], 16),
                int(primary_hex[2:4], 16),
                int(primary_hex[4:6], 16)
            )
        except:
            primary_color = RgbColor(0, 102, 204)

        # Generate slides
        for slide_data in content.get("slides", []):
            slide_type = slide_data.get("type", "content")

            if slide_type == "title":
                self._add_title_slide(prs, slide_data, primary_color)
            elif slide_type == "agenda":
                self._add_agenda_slide(prs, slide_data, primary_color)
            elif slide_type == "content":
                self._add_content_slide(prs, slide_data, primary_color)
            elif slide_type == "two_column":
                self._add_two_column_slide(prs, slide_data, primary_color)
            elif slide_type == "quote":
                self._add_quote_slide(prs, slide_data, primary_color)
            elif slide_type == "closing":
                self._add_closing_slide(prs, slide_data, primary_color)
            else:
                self._add_content_slide(prs, slide_data, primary_color)

        # Save to bytes
        pptx_buffer = io.BytesIO()
        prs.save(pptx_buffer)
        pptx_buffer.seek(0)
        return pptx_buffer.getvalue()

    def _add_title_slide(self, prs, data: Dict[str, Any], primary_color):
        """Add title slide."""
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN

        slide_layout = prs.slide_layouts[6]  # Blank
        slide = prs.slides.add_slide(slide_layout)

        # Title
        title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(1.5))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = data.get("title", "Title")
        p.font.size = Pt(54)
        p.font.bold = True
        p.font.color.rgb = primary_color
        p.alignment = PP_ALIGN.CENTER

        # Subtitle
        if data.get("subtitle"):
            subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.333), Inches(1))
            tf = subtitle_box.text_frame
            p = tf.paragraphs[0]
            p.text = data.get("subtitle", "")
            p.font.size = Pt(24)
            p.alignment = PP_ALIGN.CENTER

    def _add_agenda_slide(self, prs, data: Dict[str, Any], primary_color):
        """Add agenda slide."""
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN

        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = data.get("title", "Agenda")
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = primary_color

        # Items
        content_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(10), Inches(5))
        tf = content_box.text_frame
        for i, item in enumerate(data.get("items", [])):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = f"  {i + 1}.  {item}"
            p.font.size = Pt(28)
            p.space_after = Pt(20)

    def _add_content_slide(self, prs, data: Dict[str, Any], primary_color):
        """Add content slide."""
        from pptx.util import Inches, Pt

        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = data.get("title", "Content")
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = primary_color

        # Content
        y_pos = 1.8
        if data.get("content"):
            content_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(12.333), Inches(1))
            tf = content_box.text_frame
            p = tf.paragraphs[0]
            p.text = data.get("content", "")
            p.font.size = Pt(20)
            y_pos += 1

        # Bullets
        if data.get("bullets"):
            bullets_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(12.333), Inches(4))
            tf = bullets_box.text_frame
            for i, bullet in enumerate(data.get("bullets", [])):
                p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                p.text = f"•  {bullet}"
                p.font.size = Pt(22)
                p.space_after = Pt(14)

    def _add_two_column_slide(self, prs, data: Dict[str, Any], primary_color):
        """Add two-column slide."""
        from pptx.util import Inches, Pt

        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12.333), Inches(1))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = data.get("title", "Comparison")
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = primary_color

        # Left column
        left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(5.8), Inches(5))
        tf = left_box.text_frame
        p = tf.paragraphs[0]
        p.text = data.get("left_title", "Left")
        p.font.size = Pt(24)
        p.font.bold = True
        for item in data.get("left_content", []):
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(18)

        # Right column
        right_box = slide.shapes.add_textbox(Inches(7), Inches(1.8), Inches(5.8), Inches(5))
        tf = right_box.text_frame
        p = tf.paragraphs[0]
        p.text = data.get("right_title", "Right")
        p.font.size = Pt(24)
        p.font.bold = True
        for item in data.get("right_content", []):
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(18)

    def _add_quote_slide(self, prs, data: Dict[str, Any], primary_color):
        """Add quote slide."""
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN

        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Quote
        quote_box = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10.333), Inches(2))
        tf = quote_box.text_frame
        p = tf.paragraphs[0]
        p.text = f'"{data.get("quote", "Quote")}"'
        p.font.size = Pt(32)
        p.font.italic = True
        p.alignment = PP_ALIGN.CENTER

        # Author
        if data.get("author"):
            author_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.8), Inches(10.333), Inches(0.5))
            tf = author_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"— {data.get('author', '')}"
            p.font.size = Pt(20)
            p.font.color.rgb = primary_color
            p.alignment = PP_ALIGN.CENTER

    def _add_closing_slide(self, prs, data: Dict[str, Any], primary_color):
        """Add closing slide."""
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN

        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Title
        title_box = slide.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11.333), Inches(1.5))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = data.get("title", "Thank You")
        p.font.size = Pt(54)
        p.font.bold = True
        p.font.color.rgb = primary_color
        p.alignment = PP_ALIGN.CENTER

        # Subtitle
        if data.get("subtitle"):
            subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(11.333), Inches(1))
            tf = subtitle_box.text_frame
            p = tf.paragraphs[0]
            p.text = data.get("subtitle", "")
            p.font.size = Pt(24)
            p.alignment = PP_ALIGN.CENTER

    def _create_placeholder_pptx(self, content: Dict[str, Any]) -> bytes:
        """Create a minimal PPTX placeholder when python-pptx is not available."""
        import zipfile

        buffer = io.BytesIO()

        # Create minimal PPTX structure
        with zipfile.ZipFile(buffer, "w") as zf:
            # Content Types
            content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
</Types>"""
            zf.writestr("[Content_Types].xml", content_types)

            # Relationships
            rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
</Relationships>"""
            zf.writestr("_rels/.rels", rels)

            # Minimal presentation
            presentation = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:sldIdLst/>
</p:presentation>"""
            zf.writestr("ppt/presentation.xml", presentation)

        buffer.seek(0)
        return buffer.getvalue()
