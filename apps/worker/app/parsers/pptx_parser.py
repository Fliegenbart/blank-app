import io
import zipfile
from typing import Optional, Dict, List
from xml.etree import ElementTree as ET
import structlog

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR

from app.parsers.base import (
    DocumentParser,
    ParsedDocument,
    ParsedPage,
    TextRun,
    ExtractedImage,
    ThemeData,
)

logger = structlog.get_logger()

# XML namespaces for OOXML
NAMESPACES = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


class PPTXParser(DocumentParser):
    """Parser for PowerPoint (PPTX) files."""

    def can_parse(self, file_type: str) -> bool:
        return file_type.lower() in ("pptx", "ppt")

    def parse(self, data: bytes) -> ParsedDocument:
        """Parse a PPTX file."""
        logger.info("Parsing PPTX document")

        doc = ParsedDocument(filename="", file_type="pptx")

        # Parse using python-pptx
        try:
            prs = Presentation(io.BytesIO(data))
            doc = self._parse_with_python_pptx(prs, doc)
        except Exception as e:
            logger.error("Error parsing with python-pptx", error=str(e))

        # Also parse raw OOXML for theme data
        try:
            theme_data = self._parse_ooxml_theme(data)
            if theme_data:
                doc.theme = theme_data
        except Exception as e:
            logger.warning("Error parsing OOXML theme", error=str(e))

        # Extract embedded images
        try:
            images = self._extract_images(data)
            if doc.pages and images:
                # Distribute images across pages (simplified)
                for i, img in enumerate(images):
                    page_idx = i % len(doc.pages)
                    doc.pages[page_idx].images.append(img)
        except Exception as e:
            logger.warning("Error extracting images", error=str(e))

        logger.info("PPTX parsing complete", pages=len(doc.pages))
        return doc

    def _parse_with_python_pptx(self, prs: Presentation, doc: ParsedDocument) -> ParsedDocument:
        """Parse presentation using python-pptx."""
        for slide_idx, slide in enumerate(prs.slides):
            page = ParsedPage(page_number=slide_idx + 1)

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            text_run = TextRun(
                                text=run.text,
                                font_name=self._get_font_name(run),
                                font_size=self._get_font_size(run),
                                font_color=self._get_font_color(run),
                                bold=run.font.bold or False,
                                italic=run.font.italic or False,
                            )
                            if text_run.text.strip():
                                page.text_runs.append(text_run)

            doc.pages.append(page)

        return doc

    def _get_font_name(self, run) -> Optional[str]:
        """Get font name from a run."""
        try:
            if run.font.name:
                return run.font.name
        except Exception:
            pass
        return None

    def _get_font_size(self, run) -> Optional[float]:
        """Get font size in points from a run."""
        try:
            if run.font.size:
                return run.font.size.pt
        except Exception:
            pass
        return None

    def _get_font_color(self, run) -> Optional[str]:
        """Get font color as hex from a run."""
        try:
            color = run.font.color
            if color and color.rgb:
                return f"#{color.rgb}"
            elif color and color.theme_color:
                # Return theme color name for now
                return f"theme:{color.theme_color}"
        except Exception:
            pass
        return None

    def _parse_ooxml_theme(self, data: bytes) -> Optional[ThemeData]:
        """Parse theme data from raw OOXML."""
        theme = ThemeData()

        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            # Find theme file
            theme_files = [n for n in zf.namelist() if "theme" in n.lower() and n.endswith(".xml")]

            for theme_file in theme_files:
                try:
                    content = zf.read(theme_file)
                    root = ET.fromstring(content)

                    # Parse color scheme
                    for clr_scheme in root.iter(f"{{{NAMESPACES['a']}}}clrScheme"):
                        for child in clr_scheme:
                            color_name = child.tag.split("}")[-1]
                            # Look for srgbClr or sysClr
                            srgb = child.find(f".//{{{NAMESPACES['a']}}}srgbClr")
                            if srgb is not None:
                                theme.color_scheme[color_name] = f"#{srgb.get('val')}"
                            sys_clr = child.find(f".//{{{NAMESPACES['a']}}}sysClr")
                            if sys_clr is not None and sys_clr.get("lastClr"):
                                theme.color_scheme[color_name] = f"#{sys_clr.get('lastClr')}"

                    # Parse font scheme
                    for font_scheme in root.iter(f"{{{NAMESPACES['a']}}}fontScheme"):
                        major_font = font_scheme.find(f".//{{{NAMESPACES['a']}}}majorFont")
                        minor_font = font_scheme.find(f".//{{{NAMESPACES['a']}}}minorFont")

                        if major_font is not None:
                            latin = major_font.find(f"{{{NAMESPACES['a']}}}latin")
                            if latin is not None:
                                theme.font_scheme["heading"] = latin.get("typeface")

                        if minor_font is not None:
                            latin = minor_font.find(f"{{{NAMESPACES['a']}}}latin")
                            if latin is not None:
                                theme.font_scheme["body"] = latin.get("typeface")

                except Exception as e:
                    logger.warning("Error parsing theme file", file=theme_file, error=str(e))

        # Extract accent colors
        accent_names = ["accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]
        theme.accent_colors = [
            theme.color_scheme.get(name, "")
            for name in accent_names
            if theme.color_scheme.get(name)
        ]

        return theme if theme.color_scheme or theme.font_scheme else None

    def _extract_images(self, data: bytes) -> List[ExtractedImage]:
        """Extract embedded images from PPTX."""
        images = []

        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            # Look for images in ppt/media/
            for name in zf.namelist():
                if name.startswith("ppt/media/") and any(
                    name.lower().endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".gif", ".bmp"]
                ):
                    try:
                        img_data = zf.read(name)
                        from PIL import Image
                        img = Image.open(io.BytesIO(img_data))
                        width, height = img.size
                        img_format = img.format.lower() if img.format else "png"

                        images.append(
                            ExtractedImage(
                                data=img_data,
                                format=img_format,
                                width=width,
                                height=height,
                                filename=name.split("/")[-1],
                            )
                        )
                    except Exception as e:
                        logger.warning("Error extracting image", name=name, error=str(e))

        return images
